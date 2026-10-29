"""
revise_doc.py — additive-only revision of an approved Word document.

Used by the /upgrade-doc skill when a system moves to a new release (for
example 1.0 to 2.0). The approved document is opened as a new draft and new
content is ADDED to it. Approved content is never removed or reworded, and the
revision history table is updated. A verify step proves the approved content
is intact.

SUBCOMMANDS
-----------
  init   <approved.docx> --version 2.0 [--old-version 1.0] [--out-dir DIR]
         Copy the approved document to a new draft named for the new version
         and update the version label in headers, footers, and the first
         15 body paragraphs. Prints the draft path as JSON.

  add    <draft.docx> <additions.json> [--highlight]
         additions.json: { "<exact heading text>": ["paragraph", ...] }
         Each paragraph list is inserted at the END of that section, in the
         formatting of the section's last body paragraph. An item may also be
         { "text": "...", "style": "<paragraph style name>" }.
         Headings not found in the document are reported, never created.

  revlog <draft.docx> --version V --date D --summary S [--by B] [--ref R]
         Add a row to the revision history table (the table whose header
         row holds Version and Summary/Change columns).

  verify <approved.docx> <draft.docx> [--old-version 1.0] [--new-version 2.0]
         Check that every paragraph and table cell of the approved document
         is still present, unchanged and in order, and that every approved
         revision history row is still present. Exit code 1 on any violation.

Dependencies: python-docx (see requirements.txt).
"""

import argparse
import copy
import difflib
import json
import re
import shutil
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

sys.path.insert(0, str(Path(__file__).resolve().parent))
from generate_doc import _is_heading_style  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUT = REPO_ROOT / "deliverables" / "in-progress"
COVER_PARAGRAPHS = 15

VERSION_WORD = r"(?i)(\bv(?:ersion)?\s*\.?\s*:?\s*)"


def _version_re(version):
    return re.compile(VERSION_WORD + re.escape(version) + r"(?![\d.])")


def _body_items(doc):
    """Yield ('p', Paragraph) / ('t', Table) in document order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield "p", Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield "t", Table(child, doc)


# --- revision history table -------------------------------------------------

def _header_cells(table):
    return [c.text.strip().lower() for c in table.rows[0].cells]


def find_revision_table(doc):
    for table in doc.tables:
        head = _header_cells(table)
        if (any("version" in h for h in head)
                and any(("summary" in h or "change" in h) for h in head)):
            return table
    return None


def _column_map(table):
    head = _header_cells(table)
    mapping = {}
    for i, h in enumerate(head):
        if "version" in h:
            mapping["version"] = i
        elif "date" in h:
            mapping["date"] = i
        elif "summary" in h or "change" in h:
            mapping["summary"] = i
        elif "by" in h:
            mapping["by"] = i
        elif "reference" in h:
            mapping["ref"] = i
    return mapping


def _is_blank_row(texts):
    """True for an empty row or an unfilled template row of <placeholder> cells."""
    return all(not t.strip() or re.fullmatch(r"<[^<>]*>(\s+or\s+N/A)?", " ".join(t.split()))
               for t in texts)


def _set_cell(cell, text):
    first = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    if first.runs:
        first.runs[0].text = text
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.add_run(text)


# --- init -------------------------------------------------------------------

def _bump_paragraph(par, pattern, new):
    """Replace the version label when it sits inside a single run."""
    hit = False
    for run in par.runs:
        if pattern.search(run.text):
            run.text = pattern.sub(lambda m: m.group(1) + new, run.text)
            hit = True
    return hit or not pattern.search(par.text)


def cmd_init(args):
    from docx import Document
    src = Path(args.approved)
    if not src.is_file():
        print("Error: {} not found.".format(src))
        return 1
    old = args.old_version
    if not old:
        m = re.search(r"v?(\d+\.\d+)", src.stem)
        if not m:
            print("Error: cannot read the current version from the filename; "
                  "pass --old-version.")
            return 1
        old = m.group(1)
    stem = src.stem
    if re.search(r"v?" + re.escape(old), stem):
        stem = re.sub(r"(v?)" + re.escape(old), lambda m: m.group(1) + args.version,
                      stem, count=1)
    else:
        stem = "{}_v{}".format(stem, args.version)
    if not stem.lower().endswith("_draft"):
        stem += "_draft"
    out_dir = Path(args.out_dir) if args.out_dir else DEFAULT_OUT
    out_dir.mkdir(parents=True, exist_ok=True)
    dest = out_dir / (stem + src.suffix)
    if dest.exists():
        print("Error: {} already exists; refusing to overwrite.".format(dest))
        return 1
    shutil.copy2(src, dest)

    doc = Document(str(dest))
    pattern = _version_re(old)
    unresolved = []
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer):
            if part.is_linked_to_previous:
                continue
            for par in part.paragraphs:
                if not _bump_paragraph(par, pattern, args.version):
                    unresolved.append(par.text.strip())
    for par in doc.paragraphs[:COVER_PARAGRAPHS]:
        if not _bump_paragraph(par, pattern, args.version):
            unresolved.append(par.text.strip())
    doc.save(str(dest))

    other = [p.text.strip()[:80] for p in doc.paragraphs[COVER_PARAGRAPHS:]
             if pattern.search(p.text)]
    print(json.dumps({
        "draft": dest.relative_to(REPO_ROOT).as_posix()
        if REPO_ROOT in dest.parents else str(dest),
        "old_version": old, "new_version": args.version,
        "version_label_not_updated": unresolved,
        "old_version_mentions_in_body_left_unchanged": other,
        "revision_table_found": find_revision_table(doc) is not None,
    }, indent=2, ensure_ascii=False))
    return 0


# --- add --------------------------------------------------------------------

def _section_end(items, start):
    """Index of the last body item belonging to the section opened at `start`."""
    end = start
    for j in range(start + 1, len(items)):
        kind, obj = items[j]
        if kind == "p" and _is_heading_style(obj.style) and obj.text.strip():
            break
        end = j
    return end


def _new_paragraph_after(anchor_el, template_par, text, style_name, doc, highlight):
    from docx.enum.text import WD_COLOR_INDEX
    from docx.text.paragraph import Paragraph
    new_el = copy.deepcopy(template_par._element)
    anchor_el.addnext(new_el)
    par = Paragraph(new_el, template_par._parent)
    for run in par.runs[1:]:
        run._element.getparent().remove(run._element)
    for extra in list(new_el):
        if extra.tag.endswith(("}commentRangeStart", "}commentRangeEnd", "}bookmarkStart",
                               "}bookmarkEnd", "}hyperlink")):
            new_el.remove(extra)
    if par.runs:
        par.runs[0].text = text
        run = par.runs[0]
    else:
        run = par.add_run(text)
    if style_name:
        par.style = doc.styles[style_name]
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return par


def cmd_add(args):
    from docx import Document
    path = Path(args.draft)
    additions = json.loads(Path(args.additions).read_text(encoding="utf-8-sig"))
    doc = Document(str(path))
    items = list(_body_items(doc))
    added, not_found = {}, []
    for heading, paras in additions.items():
        if isinstance(paras, str):
            paras = [paras]
        start = next((i for i, (k, o) in enumerate(items)
                      if k == "p" and _is_heading_style(o.style)
                      and o.text.strip() == heading.strip()), None)
        if start is None:
            not_found.append(heading)
            continue
        end = _section_end(items, start)
        body = [o for k, o in items[start + 1:end + 1] if k == "p" and o.text.strip()]
        template = body[-1] if body else None
        if template is None:
            not_found.append(heading + " (section has no body paragraph to copy formatting from)")
            continue
        anchor = items[end][1]._element if items[end][0] == "p" else items[end][1]._tbl
        for item in paras:
            text = item["text"] if isinstance(item, dict) else item
            style = item.get("style") if isinstance(item, dict) else None
            par = _new_paragraph_after(anchor, template, text, style, doc, args.highlight)
            anchor = par._element
        added[heading] = len(paras)
    doc.save(str(path))
    print(json.dumps({"added_paragraphs_by_section": added,
                      "headings_not_found": not_found}, indent=2, ensure_ascii=False))
    return 1 if not_found else 0


# --- revlog -----------------------------------------------------------------

def cmd_revlog(args):
    from docx import Document
    from docx.table import _Row
    path = Path(args.draft)
    doc = Document(str(path))
    table = find_revision_table(doc)
    if table is None:
        print("Error: no revision history table found (a table whose header has "
              "Version and Summary/Change columns). Add the entry by hand.")
        return 1
    cols = _column_map(table)
    values = {"version": args.version, "date": args.date, "summary": args.summary,
              "by": args.by, "ref": args.ref}
    last = table.rows[-1]
    blank = len(table.rows) > 1 and _is_blank_row([c.text for c in last.cells])
    if blank:
        row_cells = last.cells
    else:
        new_tr = copy.deepcopy(last._tr)
        last._tr.addnext(new_tr)
        row_cells = _Row(new_tr, table).cells
        for cell in row_cells:
            _set_cell(cell, "")
    for key, idx in cols.items():
        if values.get(key):
            _set_cell(row_cells[idx], values[key])
    doc.save(str(path))
    print("Revision history updated: version {}.".format(args.version))
    return 0


# --- verify -----------------------------------------------------------------

def _units(doc, version_pattern):
    """Ordered list of normalised paragraph / cell texts, excluding the
    revision history table, plus that table's rows."""
    rev = find_revision_table(doc)
    rev_el = rev._tbl if rev is not None else None
    units = []

    def norm(text):
        text = " ".join(text.split())
        return version_pattern.sub(lambda m: m.group(1) + "{VERSION}", text) \
            if version_pattern else text

    for kind, obj in _body_items(doc):
        if kind == "p":
            if obj.text.strip():
                units.append(norm(obj.text))
        elif obj._tbl is not rev_el:
            for row in obj.rows:
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    if cell.text.strip():
                        units.append(norm(cell.text))
    rows = []
    if rev is not None:
        rows = [tuple(" ".join(c.text.split()) for c in r.cells) for r in rev.rows[1:]]
        rows = [r for r in rows if not _is_blank_row(r)]
    return units, rows


def cmd_verify(args):
    from docx import Document
    old = args.old_version
    if not old:
        m = re.search(r"v?(\d+\.\d+)", Path(args.approved).stem)
        old = m.group(1) if m else None
    pat = _version_re(old) if old else None
    a_units, a_rows = _units(Document(args.approved), pat)
    b_units, b_rows = _units(Document(args.draft), pat)

    violations = []
    matcher = difflib.SequenceMatcher(None, a_units, b_units, autojunk=False)
    added = 0
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "insert":
            added += j2 - j1
        elif tag == "delete":
            for t in a_units[i1:i2]:
                violations.append({"type": "removed", "approved_text": t[:200]})
        elif tag == "replace":
            for t in a_units[i1:i2]:
                violations.append({"type": "modified_or_removed", "approved_text": t[:200],
                                   "draft_text_nearby": " | ".join(b_units[j1:j2])[:200]})
    for r in (r for r in a_rows if r not in b_rows):
        violations.append({"type": "revision_history_row_lost", "approved_text": " | ".join(r)})
    result = {"passed": not violations, "approved_units_checked": len(a_units),
              "units_added_in_draft": added,
              "revision_history_rows_added": max(len(b_rows) - len(a_rows), 0),
              "violations": violations}
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0 if not violations else 1


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("init")
    p.add_argument("approved")
    p.add_argument("--version", required=True)
    p.add_argument("--old-version")
    p.add_argument("--out-dir")
    p.set_defaults(fn=cmd_init)

    p = sub.add_parser("add")
    p.add_argument("draft")
    p.add_argument("additions")
    p.add_argument("--highlight", action="store_true")
    p.set_defaults(fn=cmd_add)

    p = sub.add_parser("revlog")
    p.add_argument("draft")
    p.add_argument("--version", required=True)
    p.add_argument("--date", required=True)
    p.add_argument("--summary", required=True)
    p.add_argument("--by", default="")
    p.add_argument("--ref", default="")
    p.set_defaults(fn=cmd_revlog)

    p = sub.add_parser("verify")
    p.add_argument("approved")
    p.add_argument("draft")
    p.add_argument("--old-version")
    p.add_argument("--new-version")
    p.set_defaults(fn=cmd_verify)

    args = parser.parse_args()
    return args.fn(args)


if __name__ == "__main__":
    sys.exit(main())
