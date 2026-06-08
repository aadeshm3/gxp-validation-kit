"""
check_status.py — Report status of all validation deliverables and open items.

PURPOSE
-------
Reconciles the project state recorded in MASTER_CONTEXT.md and
DELIVERABLE_STATUS.md against the actual files on disk, and writes a markdown
status report to stdout.

INPUTS
------
- MASTER_CONTEXT.md (repo root): Last Refreshed date (header), section 6
  (Open Items & Pending Confirmations).
- DELIVERABLE_STATUS.md (repo root): markdown table of recorded deliverables.
- deliverables/in-progress/ and deliverables/approved/: actual .docx and .md files.
- context/ subtree: files newer than Last Refreshed are "unprocessed".

OUTPUT
------
A markdown report on stdout with four sections: DELIVERABLES, OPEN ITEMS,
UNPROCESSED FILES IN CONTEXT, and a single NEXT RECOMMENDED ACTION line.

The script never raises and never exits non-zero on missing optional files;
missing files are treated as empty.

DEPENDENCIES
------------
Standard library only. python-docx is used opportunistically for .docx
placeholder counting if available; absence is handled gracefully.

INTEGRATION
-----------
Backend for the /check-status skill and the status-report.yml workflow.
Read-only — never modifies project state.
"""

import os
import re
import sys
from datetime import datetime
from pathlib import Path

# Ensure UTF-8 output so em-dashes and other characters render on every platform.
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

REPO_ROOT = Path(__file__).resolve().parent.parent
CONTEXT_DIR = REPO_ROOT / "context"
IN_PROGRESS_DIR = REPO_ROOT / "deliverables" / "in-progress"
APPROVED_DIR = REPO_ROOT / "deliverables" / "approved"
CONTEXT_FILE = REPO_ROOT / "MASTER_CONTEXT.md"
STATUS_FILE = REPO_ROOT / "DELIVERABLE_STATUS.md"

PLACEHOLDER_PATTERN = "[CONFIRM:"


def read_text(path):
    """Return file text, or empty string if the file is missing or unreadable."""
    try:
        return path.read_text(encoding="utf-8", errors="replace")
    except (OSError, UnicodeError):
        return ""


def parse_last_refreshed(master_text):
    """Extract the Last Refreshed date from the MASTER_CONTEXT header.

    Returns a datetime, or None if not found or not parseable.
    """
    match = re.search(r"Last Refreshed:\s*\[?([0-9]{4}-[0-9]{2}-[0-9]{2})\]?", master_text)
    if not match:
        return None
    try:
        return datetime.strptime(match.group(1), "%Y-%m-%d")
    except ValueError:
        return None


def parse_markdown_table(text, header_must_contain):
    """Parse the first markdown table whose header row contains the given token.

    Returns a list of row-cell lists (data rows only, separator row dropped).
    """
    rows = []
    in_table = False
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped.startswith("|"):
            if in_table:
                break
            continue
        cells = [c.strip() for c in stripped.strip("|").split("|")]
        if not in_table:
            if header_must_contain.lower() in stripped.lower():
                in_table = True
            continue
        if set("".join(cells)) <= set("-: "):
            continue  # separator row
        rows.append(cells)
    return rows


def parse_open_items(master_text):
    """Return data rows of the section 6 Open Items & Pending Confirmations table."""
    section = ""
    capture = False
    for line in master_text.splitlines():
        if line.strip().startswith("## 6"):
            capture = True
            continue
        if capture and line.strip().startswith("## "):
            break
        if capture:
            section += line + "\n"
    return parse_markdown_table(section, "Description")


def count_placeholders(path):
    """Count [CONFIRM: occurrences in a deliverable file (.md or .docx)."""
    if path.suffix.lower() == ".md":
        return read_text(path).count(PLACEHOLDER_PATTERN)
    if path.suffix.lower() == ".docx":
        try:
            from docx import Document
        except ImportError:
            return None
        try:
            doc = Document(str(path))
        except Exception:
            return None
        text = "\n".join(p.text for p in doc.paragraphs)
        return text.count(PLACEHOLDER_PATTERN)
    return 0


def scan_deliverables(directory):
    """Return list of (path, placeholder_count) for .md/.docx files in a directory."""
    results = []
    if not directory.is_dir():
        return results
    for path in sorted(directory.iterdir()):
        if path.suffix.lower() in (".md", ".docx") and path.name != ".gitkeep":
            results.append((path, count_placeholders(path)))
    return results


def find_unprocessed(last_refreshed):
    """Return context/ files modified after last_refreshed (all files if None)."""
    unprocessed = []
    if not CONTEXT_DIR.is_dir():
        return unprocessed
    for path in CONTEXT_DIR.rglob("*"):
        if not path.is_file() or path.name == ".gitkeep":
            continue
        if last_refreshed is None:
            unprocessed.append(path)
            continue
        mtime = datetime.fromtimestamp(os.path.getmtime(path))
        if mtime > last_refreshed:
            unprocessed.append(path)
    return unprocessed


def rel(path):
    """Repo-relative POSIX path for display."""
    try:
        return path.relative_to(REPO_ROOT).as_posix()
    except ValueError:
        return path.as_posix()


def main():
    master_text = read_text(CONTEXT_FILE)
    status_text = read_text(STATUS_FILE)
    last_refreshed = parse_last_refreshed(master_text)

    in_progress = scan_deliverables(IN_PROGRESS_DIR)
    approved = scan_deliverables(APPROVED_DIR)
    recorded = parse_markdown_table(status_text, "Document")
    open_items = parse_open_items(master_text)
    unprocessed = find_unprocessed(last_refreshed)

    on_disk_names = {p.name for p, _ in in_progress + approved}
    recorded_locations = " ".join(" ".join(r) for r in recorded)

    lines = []
    lines.append("# Status Report")
    lines.append("")
    lines.append("Generated: " + datetime.now().strftime("%Y-%m-%d %H:%M"))
    lines.append("Last Refreshed (MASTER_CONTEXT): " +
                 (last_refreshed.strftime("%Y-%m-%d") if last_refreshed else "not set"))
    lines.append("")

    lines.append("## Deliverables")
    lines.append("")
    lines.append("| Document | Location | Placeholders remaining | Next action |")
    lines.append("|----------|----------|------------------------|-------------|")
    if not in_progress and not approved:
        lines.append("| None found on disk | — | — | Run /generate-doc to create the first deliverable |")
    for path, count in in_progress:
        ph = "unknown (python-docx not installed)" if count is None else str(count)
        action = "Resolve placeholders" if count else "Run /gap-check before routing"
        lines.append("| {} | {} | {} | {} |".format(path.name, rel(path), ph, action))
    for path, count in approved:
        ph = "unknown (python-docx not installed)" if count is None else str(count)
        lines.append("| {} | {} | {} | Approved |".format(path.name, rel(path), ph))

    # Discrepancies: recorded but not on disk.
    missing = []
    for row in recorded:
        joined = " ".join(row)
        filename_tokens = re.findall(r"[\w\-.]+\.(?:md|docx)", joined)
        for token in filename_tokens:
            if token not in on_disk_names:
                missing.append(token)
    if missing:
        lines.append("")
        lines.append("Discrepancies — recorded in DELIVERABLE_STATUS.md but not found on disk: " +
                     ", ".join(sorted(set(missing))))
    orphan = [p.name for p, _ in in_progress + approved if p.name not in recorded_locations]
    if orphan:
        lines.append("")
        lines.append("Discrepancies — on disk but not recorded in DELIVERABLE_STATUS.md: " +
                     ", ".join(sorted(set(orphan))))

    lines.append("")
    lines.append("## Open Items")
    lines.append("")
    if open_items:
        lines.append("| " + " | ".join(open_items[0]) + " |")
        lines.append("|" + "|".join(["---"] * len(open_items[0])) + "|")
        for row in open_items[1:]:
            lines.append("| " + " | ".join(row) + " |")
    else:
        lines.append("No open items recorded in MASTER_CONTEXT.md section 6.")

    lines.append("")
    lines.append("## Unprocessed Files in context/")
    lines.append("")
    if unprocessed:
        for path in unprocessed:
            lines.append("- " + rel(path))
    else:
        lines.append("None. All context files are reflected in the last refresh.")

    # Single recommended next action.
    if not master_text or "[populated by /build-context]" in master_text:
        action = "Run /build-context — MASTER_CONTEXT.md is not yet populated."
    elif unprocessed:
        action = "Run /update-context — {} context file(s) changed since the last refresh.".format(len(unprocessed))
    elif any(c for _, c in in_progress if c):
        action = "Resolve open [CONFIRM:] placeholders, then run /gap-check on affected drafts."
    elif in_progress:
        action = "Run /gap-check on in-progress drafts before routing for review."
    else:
        action = "No action required. All recorded deliverables are reconciled."

    lines.append("")
    lines.append("## Next Recommended Action")
    lines.append("")
    lines.append(action)
    lines.append("")

    print("\n".join(lines))


if __name__ == "__main__":
    main()
